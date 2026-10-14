import json

import openpyxl

from rag_harness import export


def test_export_builds_one_row_per_run_with_detail_sheets(tmp_path, monkeypatch):
    man, raw = tmp_path / "manifests", tmp_path / "raw"
    man.mkdir()
    rid = "E1-scifact-c2-r1-abc123"
    cfg = {"dataset": {"name": "scifact"}, "ingest": {"concurrency": 2},
           "embedding": {"workers_per_invocation": 8, "dimensions": 1024}, "lambdas": {"embed": {"memory_mb": 1769}}}
    (man / f"{rid}.json").write_text(json.dumps({
        "run_id": rid, "experiment_id": "E1", "repetition": 1, "config": cfg,
        "phases": {"deploy_end": "2026-10-14T00:00:00+00:00"},
        "ingest_expected": {"docs": 2}, "ingest_result": {"vectors_expected": 4, "vectors_stored": 4, "dlq_messages": 0},
        "cost_estimate_usd": {"titan": 0.001, "lambda": 0.002}}))
    (raw / rid).mkdir(parents=True)
    (raw / rid / "embed_saaf.jsonl").write_text("\n".join(json.dumps(
        {"run_id": rid, "startTime": 1000 + 500 * k, "endTime": 1400 + 500 * k, "runtime": 400, "chunks": 2}) for k in range(2)))
    (raw / rid / "client_queries.jsonl").write_text(json.dumps({"i": 0, "qid": "q1", "warmup": False, "status": "ok", "rtt_ms": 500}))
    (raw / rid / "query_saaf.jsonl").write_text(json.dumps({"qid": "q1", "handler_ms": 480, "gen_ttft_ms": 250}))
    monkeypatch.setattr(export, "MANIFESTS", man)
    monkeypatch.setattr(export, "RAW", raw)
    wb = openpyxl.load_workbook(export.build(tmp_path / "out.xlsx"))
    runs = wb["Runs"]
    head = [c.value for c in runs[1]]
    row = dict(zip(head, [c.value for c in runs[2]]))
    assert row["run_id"] == rid and row["concurrency"] == 2 and row["ingest_s"] == 0.9
    assert row["embed_invocations"].startswith("=COUNTIFS(Embed!")
    assert wb["Embed"].max_row == 3 and wb["Queries"].max_row == 2
    assert wb["Queries"]["F2"].value == 500 and wb["Queries"]["G2"].value == 480
